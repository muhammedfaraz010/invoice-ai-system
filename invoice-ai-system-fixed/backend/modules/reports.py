"""
Invoice Amount & Tax Report Generator
--------------------------------------
Builds a full invoice amount report (Excel or PDF) using the country-aware
tax metadata from utils.tax_display. One row per invoice, plus a summary
sheet/section with totals and a list of amount discrepancies.

This module does not touch OCR, extraction, or the database schema — it
only reads already-stored Invoice fields and formats them for export.
"""

import io
from datetime import datetime
from typing import List, Dict, Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from fpdf import FPDF

from utils.tax_display import build_tax_metadata


def _invoice_to_report_row(invoice) -> Dict[str, Any]:
    tax_meta = build_tax_metadata(
        currency=invoice.currency,
        subtotal=invoice.subtotal,
        tax_amount=invoice.tax_amount,
        total_amount=invoice.total_amount,
        vendor_gstin=invoice.vendor_gstin,
        buyer_gstin=invoice.buyer_gstin,
        vendor_vat=invoice.vendor_vat,
        invoice_date=invoice.invoice_date,
    )
    amount_check = tax_meta["amount_check"]

    return {
        "invoice_number": invoice.invoice_number or "-",
        "vendor_name": invoice.vendor_name or "-",
        "invoice_date": invoice.invoice_date or "-",
        "currency": invoice.currency or "-",
        "country": tax_meta["country"] or "Unknown",
        "subtotal": invoice.subtotal,
        "tax_name": tax_meta["tax_name"],
        "tax_display": tax_meta["tax_display"],
        "tax_amount": invoice.tax_amount,
        "total_amount": invoice.total_amount,
        "tax_registration_number": tax_meta["tax_registration_number"] or "-",
        "validation_status": invoice.validation_status or "-",
        "expected_total": amount_check["expected_total"],
        "amount_mismatch": amount_check["mismatch"],
        "difference": amount_check["difference"],
        "note": amount_check["note"],
        "warnings": "; ".join(tax_meta["warnings"]) if tax_meta["warnings"] else "",
    }


def build_report_rows(invoices) -> List[Dict[str, Any]]:
    return [_invoice_to_report_row(inv) for inv in invoices]


# ──────────────────────────────────────────────
# Excel report
# ──────────────────────────────────────────────

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_MISMATCH_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
_OK_FILL = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")

_COLUMNS = [
    ("invoice_number", "Invoice #", 16),
    ("vendor_name", "Vendor", 24),
    ("invoice_date", "Date", 14),
    ("country", "Country", 10),
    ("tax_display", "Tax", 20),
    ("tax_registration_number", "Tax Reg. No.", 20),
    ("subtotal", "Subtotal", 14),
    ("tax_amount", "Tax Amount", 14),
    ("total_amount", "Total", 14),
    ("expected_total", "Expected Total", 14),
    ("difference", "Difference", 12),
    ("validation_status", "Status", 14),
    ("warnings", "Warnings", 30),
]


def generate_excel_report(rows: List[Dict[str, Any]]) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Invoices"
    ws.append([label for _, label, _ in _COLUMNS])
    for col_idx in range(1, len(_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([row.get(key) for key, _, _ in _COLUMNS])
        row_idx = ws.max_row
        fill = _MISMATCH_FILL if row.get("amount_mismatch") else _OK_FILL
        for col_idx in range(1, len(_COLUMNS) + 1):
            ws.cell(row=row_idx, column=col_idx).fill = fill

    for idx, (_, _, width) in enumerate(_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    # Summary sheet
    summary = wb.create_sheet("Summary")
    total_invoices = len(rows)
    total_amount = sum(r["total_amount"] or 0 for r in rows)
    total_tax = sum(r["tax_amount"] or 0 for r in rows)
    mismatches = [r for r in rows if r["amount_mismatch"]]

    summary_lines = [
        ("Total Invoices", total_invoices),
        ("Total Amount (all currencies combined)", round(total_amount, 2)),
        ("Total Tax Collected (all currencies combined)", round(total_tax, 2)),
        ("Invoices With Amount Mismatch", len(mismatches)),
        ("Report Generated At (UTC)", datetime.utcnow().isoformat() + "Z"),
    ]
    for label, value in summary_lines:
        summary.append([label, value])
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 24
    for r in range(1, len(summary_lines) + 1):
        summary.cell(row=r, column=1).font = Font(bold=True)

    if mismatches:
        summary.append([])
        summary.append(["Invoices Needing Review (amount mismatch)"])
        summary.cell(row=summary.max_row, column=1).font = Font(bold=True, color="B91C1C")
        summary.append(["Invoice #", "Vendor", "Expected Total", "Actual Total", "Difference"])
        for r in range(summary.max_row - 4, summary.max_row + 1):
            pass  # header row already appended above
        for m in mismatches:
            summary.append([
                m["invoice_number"], m["vendor_name"], m["expected_total"],
                m["total_amount"], m["difference"],
            ])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ──────────────────────────────────────────────
# PDF report
# ──────────────────────────────────────────────

class _ReportPDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 10, "Invoice Amount & Tax Report", ln=True, align="C")
        self.set_font("Helvetica", "", 9)
        self.cell(0, 6, f"Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", ln=True, align="C")
        self.ln(4)


def generate_pdf_report(rows: List[Dict[str, Any]]) -> bytes:
    pdf = _ReportPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()

    total_amount = sum(r["total_amount"] or 0 for r in rows)
    total_tax = sum(r["tax_amount"] or 0 for r in rows)
    mismatches = [r for r in rows if r["amount_mismatch"]]

    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 7, f"Total Invoices: {len(rows)}    "
                   f"Total Amount: {round(total_amount, 2)}    "
                   f"Total Tax: {round(total_tax, 2)}    "
                   f"Mismatches: {len(mismatches)}", ln=True)
    pdf.ln(2)

    headers = ["Invoice #", "Vendor", "Date", "Country", "Tax", "Reg. No.",
               "Subtotal", "Tax Amt", "Total", "Status"]
    widths = [22, 35, 20, 18, 28, 28, 20, 20, 20, 22]

    pdf.set_font("Helvetica", "B", 8)
    pdf.set_fill_color(31, 41, 55)
    pdf.set_text_color(255, 255, 255)
    for h, w in zip(headers, widths):
        pdf.cell(w, 7, h, border=1, align="C", fill=True)
    pdf.ln()

    pdf.set_font("Helvetica", "", 7.5)
    pdf.set_text_color(0, 0, 0)
    for row in rows:
        fill = row.get("amount_mismatch")
        if fill:
            pdf.set_fill_color(254, 226, 226)
        else:
            pdf.set_fill_color(255, 255, 255)
        values = [
            str(row["invoice_number"])[:14], str(row["vendor_name"])[:22],
            str(row["invoice_date"])[:12], str(row["country"])[:10],
            str(row["tax_display"])[:18], str(row["tax_registration_number"])[:18],
            f"{row['subtotal']:.2f}" if row["subtotal"] is not None else "-",
            f"{row['tax_amount']:.2f}" if row["tax_amount"] is not None else "-",
            f"{row['total_amount']:.2f}" if row["total_amount"] is not None else "-",
            str(row["validation_status"])[:12],
        ]
        for v, w in zip(values, widths):
            pdf.cell(w, 6, v, border=1, fill=True)
        pdf.ln()

    if mismatches:
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(185, 28, 28)
        pdf.cell(0, 7, "Amount Discrepancy Report", ln=True)
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Helvetica", "", 8)
        for m in mismatches:
            pdf.multi_cell(
                0, 5,
                f"- Invoice {m['invoice_number']} ({m['vendor_name']}): "
                f"expected {m['expected_total']}, actual {m['total_amount']} "
                f"(difference {m['difference']:+.2f})",
            )

    return bytes(pdf.output())
